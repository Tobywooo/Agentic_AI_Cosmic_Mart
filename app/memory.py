"""Conversation memory stored as a rolling table in an Excel workbook.

Every message (customer, agent, tool call, human specialist, follow-up) is one row
tagged with its conversation_id. When the table grows past `max_rows`, the oldest
rows are dropped. The in-process copy is the source of truth while the server runs;
the workbook is rewritten after every change so it can be opened in Excel at any time.
"""

from __future__ import annotations

import logging
import os
import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

log = logging.getLogger(__name__)

COLUMNS = [
    "row_id",
    "timestamp",
    "conversation_id",
    "customer_id",
    "role",
    "tool_name",
    "content",
    "case_status",
]
COLUMN_WIDTHS = {"row_id": 8, "timestamp": 22, "conversation_id": 16, "customer_id": 12,
                 "role": 13, "tool_name": 22, "content": 100, "case_status": 13}
EXCEL_CELL_LIMIT = 32_000
SHEET_NAME = "memory"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _clean(value) -> str | int | None:
    if value is None or isinstance(value, int):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text if len(text) <= EXCEL_CELL_LIMIT else text[:EXCEL_CELL_LIMIT] + " …[truncated]"


class ExcelMemory:
    def __init__(self, path: Path, max_rows: int = 50):
        self.path = Path(path)
        self.max_rows = max_rows
        self._lock = threading.RLock()
        self._rows: list[dict] = self._load()
        self._next_id = max((r["row_id"] or 0 for r in self._rows), default=0) + 1
        self._flush()

    # ------------------------------------------------------------------ public API
    def append(
        self,
        conversation_id: str,
        role: str,
        content: str,
        *,
        customer_id: str | None = None,
        tool_name: str | None = None,
        case_status: str | None = None,
    ) -> dict:
        with self._lock:
            row = {
                "row_id": self._next_id,
                "timestamp": _now(),
                "conversation_id": conversation_id,
                "customer_id": customer_id,
                "role": role,
                "tool_name": tool_name,
                "content": _clean(content),
                "case_status": case_status,
            }
            self._next_id += 1
            self._rows.append(row)
            if len(self._rows) > self.max_rows:
                del self._rows[: len(self._rows) - self.max_rows]
            self._flush()
            return dict(row)

    def history(self, conversation_id: str) -> list[dict]:
        with self._lock:
            return [dict(r) for r in self._rows if r["conversation_id"] == conversation_id]

    def all_rows(self) -> list[dict]:
        with self._lock:
            return [dict(r) for r in self._rows]

    # ------------------------------------------------------------------ persistence
    def _load(self) -> list[dict]:
        if not self.path.exists():
            return []
        try:
            wb = load_workbook(self.path, read_only=True)
        except Exception:  # corrupt / unreadable workbook: start fresh but keep the old file
            backup = self.path.with_suffix(f".corrupt-{datetime.now():%Y%m%d%H%M%S}.xlsx")
            log.exception("Could not read %s; moving it to %s", self.path, backup)
            os.replace(self.path, backup)
            return []
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        header = [str(h) if h is not None else "" for h in rows[0]]
        records = []
        for values in rows[1:]:
            record = {col: None for col in COLUMNS}
            record.update({h: v for h, v in zip(header, values) if h in record})
            if record["conversation_id"] is None:
                continue
            record["row_id"] = int(record["row_id"]) if record["row_id"] is not None else None
            records.append(record)
        return records[-self.max_rows:]

    def _flush(self) -> None:
        """Rewrite the workbook. If Excel has it open (Windows file lock) keep going and retry next write."""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in self._rows:
            ws.append([row[c] for c in COLUMNS])
        for idx, col in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = COLUMN_WIDTHS[col]
        content_col = COLUMNS.index("content") + 1
        for (cell,) in ws.iter_rows(min_row=2, min_col=content_col, max_col=content_col):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        tmp = self.path.with_name(f"~tmp-{self.path.name}")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError:
            log.warning("%s is locked (open in Excel?). Memory kept in RAM; will retry on next write.", self.path)
            tmp.unlink(missing_ok=True)
