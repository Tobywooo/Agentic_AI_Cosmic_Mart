from openpyxl import load_workbook

from app.frustration import detect_frustration
from app.memory import ExcelMemory
from app.protocol import parse_agent_output


def test_memory_is_capped_and_survives_restart(tmp_path):
    path = tmp_path / "memory.xlsx"
    memory = ExcelMemory(path, max_rows=50)
    for i in range(60):
        memory.append("conv-a" if i % 2 else "conv-b", "user", f"message {i}")

    rows = memory.all_rows()
    assert len(rows) == 50
    assert rows[0]["content"] == "message 10" and rows[-1]["content"] == "message 59"

    ws = load_workbook(path)["memory"]
    assert ws.max_row == 51  # header + 50 rows
    assert [c.value for c in ws[1]][:3] == ["row_id", "timestamp", "conversation_id"]

    reloaded = ExcelMemory(path, max_rows=50)
    assert len(reloaded.history("conv-a")) == 25
    assert reloaded.append("conv-a", "assistant", "x")["row_id"] == 61


def test_parse_agent_output_variants():
    assert parse_agent_output('{"thought": "t", "final_answer": "Hi"}').final_answer == "Hi"
    fenced = 'Sure!\n```json\n{"action": "get_order_details", "action_input": {"order_id": "CM-1"}}\n```'
    step = parse_agent_output(fenced)
    assert step.action == "get_order_details" and step.action_input == {"order_id": "CM-1"}
    stringly = parse_agent_output('{"action": "get_order_history", "action_input": "{}"}')
    assert stringly.action == "get_order_history" and stringly.action_input == {}
    assert parse_agent_output("no json here") is None
    assert parse_agent_output('{"foo": 1}') is None


def test_frustration_detection():
    calm = detect_frustration("Hi, I'd like to return my kettle please.")
    assert calm.score == 0 and not calm.should_escalate(0.6)

    human = detect_frustration("Can I talk to a human please")
    assert human.requests_human and human.should_escalate(0.6)

    angry = detect_frustration("This is unacceptable, I'm so frustrated!!")
    assert angry.should_escalate(0.6)

    repeat = detect_frustration("My order CM-10006 has not arrived yet",
                                ["my order CM-10006 has not arrived yet."])
    assert "customer repeating themselves" in repeat.signals
